import os
import re
import time
import win32com.client
from docx import Document
from docx.shared import Pt, RGBColor
from pypdf import PdfWriter

source_dir = r"D:\LP_Agent\word_reports"
temp_docx_dir = r"D:\LP_Agent\docx_temp_perf"
temp_pdf_dir = r"D:\LP_Agent\pdf_temp_perf"

os.makedirs(temp_docx_dir, exist_ok=True)
os.makedirs(temp_pdf_dir, exist_ok=True)

files = sorted([f for f in os.listdir(source_dir) if f.endswith(".docx") and not f.startswith("~$")])

print("상단 헤더 정밀 서식(위치 굵게/큰글씨, 간격 확장, 중복 제거) 적용 중...")

for filename in files:
    doc_path = os.path.join(source_dir, filename)
    doc = Document(doc_path)
    
    # 1. 위치 번호 추출
    loc_val = ""
    for p in doc.paragraphs:
        m = re.search(r"위치\s*:\s*[@NL]?(\d+)", p.text)
        if m:
            loc_val = m.group(1)
            break
    if not loc_val:
        m_fn = re.search(r"위치_(\d+)", filename)
        if m_fn:
            loc_val = m_fn.group(1)

    # 2. 본문에서 레이블, 발매년도 데이터 추출 (음반 소개 또는 데이터 영역)
    label_val = ""
    year_val = ""
    
    full_text = "\n".join([p.text for p in doc.paragraphs])
    
    # 음반 소개 등에서 연도 추출 (예: 1983 년)
    m_yr = re.search(r"(\d{4})\s*년", full_text)
    if m_yr:
        year_val = m_yr.group(1)
        
    # 레코드사 추출
    m_lb = re.search(r"([A-Za-z가-힣0-9\s]+(?:레코드|Records|라벨|Label|모타운|Motown|아리스타|Arista))", full_text)
    if m_lb and len(m_lb.group(1).strip()) < 30:
        label_val = m_lb.group(1).strip()

    # 3. 기존의 상단 메타 문단(번호, 위치, 레이블, 발매년도) 모두 비우기
    target_idx = -1
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text
        if any(k in txt for k in ["위치 :", "위치:", "음반 레이블", "음반레이블", "발매년도", "번호 :", "번호:"]):
            if target_idx == -1:
                target_idx = idx
            p.text = ""

    # 4. 상단 헤더 문단 재구성
    if target_idx == -1:
        target_idx = 1 if len(doc.paragraphs) > 1 else 0
        p_header = doc.paragraphs[target_idx]
    else:
        p_header = doc.paragraphs[target_idx]

    p_header.text = "" # 초기화
    
    # [1] 위치 번호 (크고 굵게: 13pt, Bold)
    run_loc = p_header.add_run(f"위치 : {loc_val}")
    run_loc.bold = True
    run_loc.font.size = Pt(13)
    run_loc.font.color.rgb = RGBColor(0, 0, 0)
    
    # [2] 위치와 음반 레이블 사이 넓은 간격
    run_space1 = p_header.add_run("       |       ")
    run_space1.font.size = Pt(10.5)
    run_space1.font.color.rgb = RGBColor(120, 120, 120)
    
    # [3] 음반 레이블
    lbl_text = f"음반 레이블 : {label_val}" if label_val else "음반 레이블 : "
    run_lbl = p_header.add_run(lbl_text)
    run_lbl.font.size = Pt(10.5)
    
    # [4] 음반 레이블과 발매년도 사이 간격
    run_space2 = p_header.add_run("       |       ")
    run_space2.font.size = Pt(10.5)
    run_space2.font.color.rgb = RGBColor(120, 120, 120)
    
    # [5] 발매년도
    yr_text = f"발매년도 : {year_val}" if year_val else "발매년도 : "
    run_yr = p_header.add_run(yr_text)
    run_yr.font.size = Pt(10.5)

    temp_save_path = os.path.join(temp_docx_dir, filename)
    doc.save(temp_save_path)

print("[1단계 완료] 정밀 서식 워드 생성 완료")
print("[2단계 진행] 고속 PDF 일괄 변환 중...")

# 5. PDF 변환
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False

docx_files = sorted([f for f in os.listdir(temp_docx_dir) if f.endswith(".docx")])
total = len(docx_files)

for idx, f in enumerate(docx_files, 1):
    in_file = os.path.join(temp_docx_dir, f)
    out_file = os.path.join(temp_pdf_dir, f.replace(".docx", ".pdf"))
    
    for attempt in range(3):
        try:
            doc = word.Documents.Open(in_file, ReadOnly=True)
            doc.SaveAs(out_file, FileFormat=17)
            doc.Close(0)
            break
        except Exception:
            time.sleep(0.5)
            try: word.Quit()
            except: pass
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False

    if idx % 50 == 0 or idx == total:
        print(f"변환 진행률: {idx}/{total} ({(idx/total)*100:.1f}%)")

try:
    word.Quit()
except:
    pass

# 6. PDF 병합
print("[3단계 진행] 전체 통합 검색본 병합 중...")
merger = PdfWriter()
pdf_files = sorted([os.path.join(temp_pdf_dir, f) for f in os.listdir(temp_pdf_dir) if f.endswith(".pdf")])

for pdf in pdf_files:
    merger.append(pdf)

final_pdf_path = r"D:\LP_Agent\LP_전체음반_통합검색본.pdf"
with open(final_pdf_path, "wb") as f_out:
    merger.write(f_out)
merger.close()

# 7. 임시 정리
for f in os.listdir(temp_docx_dir):
    try: os.remove(os.path.join(temp_docx_dir, f))
    except: pass
for f in pdf_files:
    try: os.remove(f)
    except: pass

try:
    os.rmdir(temp_docx_dir)
    os.rmdir(temp_pdf_dir)
except: pass

print("="*60)
print("상단 서식 재구성 및 PDF 재생성 완료!")
print(f"최종 파일: {final_pdf_path}")
print("="*60)