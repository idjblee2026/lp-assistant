import os
import re
import glob
import time
import pandas as pd
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai

# ==========================================
# 1. 환경 설정 및 API 연결
# ==========================================
api_key = "AQ.Ab8RN6Ih8-Uo7PhIPOyNq6PwegslaIB4PGfFmuZDNoZRteVcYg"
client = genai.Client(api_key=api_key)

EXCEL_FILE = r"D:\LP_Agent\lp_list.xlsx"
WORD_DIR = r"D:\LP_Agent\word_reports"
os.makedirs(WORD_DIR, exist_ok=True)

# 한글 맑은 고딕 글꼴 적용 함수
def set_font(run, font_name="맑은 고딕", size_pt=10, bold=False, color=None):
    run.font.name = font_name
    run._r.get_or_add_rPr().set(qn('w:rFonts'), font_name)
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

# AI를 통한 음반 정보 및 해설 생성 함수
def generate_lp_content(location, artist, album):
    prompt = f"""
다음 LP 음반에 대한 정확한 발매 정보와 음악 해설을 작성해 주세요.

[음반 기본 정보]
- 위치(번호): {location}
- 아티스트: {artist}
- 앨범명: {album}

반드시 아래 형식에 맞춰 작성해 주세요:
[레이블]: (음반 발매사명)
[발매년도]: (최초 발매년도, 예: 1975)
[장르]: (음악 장르)
[수록곡]: (주요 수록곡 4~6곡 목록)
[음반 및 곡 해설]: (이 음반의 역사적 배경, 아티스트의 음악적 특징, 추천 감상 포인트 등을 3~4문단으로 품격 있고 상세하게 작성)
"""
    try:
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt
        )
        return response.text
    except Exception as e:
        print(f"AI 호출 오류 ({album}): {e}")
        return ""

# 워드 문서(.docx) 파일 생성 함수
def create_word_doc(location, artist, album, ai_text):
    doc = docx.Document()
    
    # 여백 설정
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 제목
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(f"[{location}] {artist} - {album}")
    set_font(r_title, size_pt=16, bold=True, color=RGBColor(30, 58, 138))
    p_title.paragraph_format.space_after = Pt(14)

    # AI 텍스트 파싱
    label = re.search(r"\[레이블\]\s*:\s*(.*)", ai_text)
    year = re.search(r"\[발매년도\]\s*:\s*(.*)", ai_text)
    genre = re.search(r"\[장르\]\s*:\s*(.*)", ai_text)
    tracks = re.search(r"\[수록곡\]\s*:\s*(.*)", ai_text)
    
    label_str = label.group(1).strip() if label else "미확인"
    year_str = year.group(1).strip() if year else "미확인"
    genre_str = genre.group(1).strip() if genre else "미확인"
    tracks_str = tracks.group(1).strip() if tracks else ""

    # 기본 정보 메타데이터 박스
    meta_p = doc.add_paragraph()
    set_font(meta_p.add_run(f"■ 위치 : {location}\n"), bold=True)
    set_font(meta_p.add_run(f"■ 아티스트 : {artist}\n"), bold=True)
    set_font(meta_p.add_run(f"■ 앨범명 : {album}\n"), bold=True)
    set_font(meta_p.add_run(f"■ 음반 레이블 : {label_str}\n"))
    set_font(meta_p.add_run(f"■ 발매년도 : {year_str}\n"))
    set_font(meta_p.add_run(f"■ 장르 : {genre_str}\n"))
    if tracks_str:
        set_font(meta_p.add_run(f"■ 주요 수록곡 : {tracks_str}\n"))
    meta_p.paragraph_format.space_after = Pt(16)

    # 해설 본문
    desc_match = re.search(r"\[음반 및 곡 해설\]\s*:\s*([\s\S]*)", ai_text)
    desc_text = desc_match.group(1).strip() if desc_match else ai_text

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("■ 음반 해설 및 감상 가이드")
    set_font(r_sub, size_pt=12, bold=True, color=RGBColor(31, 41, 55))
    p_sub.paragraph_format.space_after = Pt(8)

    for line in desc_text.split("\n"):
        if line.strip():
            p_desc = doc.add_paragraph()
            r_desc = p_desc.add_run(line.strip())
            set_font(r_desc, size_pt=10.5)
            p_desc.paragraph_format.line_spacing = 1.3
            p_desc.paragraph_format.space_after = Pt(6)

    # 안전한 파일명 생성 및 저장
    safe_album = re.sub(r'[\\/*?:"<>|]', "", album).strip()
    safe_artist = re.sub(r'[\\/*?:"<>|]', "", artist).strip()
    out_filename = f"위치_{location}_{safe_artist}_{safe_album}.docx"
    doc.save(os.path.join(WORD_DIR, out_filename))
    print(f"✅ 생성 완료: {out_filename}")

# ==========================================
# 2. 메인 실행 프로세스
# ==========================================
def main():
    print("==================================================")
    print("📀 LP AI 에이전트 - 신규 음반 해설 자동 생성기 📀")
    print("==================================================")

    if not os.path.exists(EXCEL_FILE):
        print(f"엑셀 파일({EXCEL_FILE})을 찾을 수 없습니다.")
        return

    df = pd.read_excel(EXCEL_FILE)
    # 컬럼명 유연 처리 (번호/위치, 아티스트/가수, 앨범명/음반명)
    col_loc = [c for c in df.columns if "위치" in str(c) or "번호" in str(c)][0]
    col_art = [c for c in df.columns if "아티스트" in str(c) or "가수" in str(c)][0]
    col_alb = [c for c in df.columns if "앨범" in str(c) or "음반" in str(c) or "제목" in str(c)][0]

    existing_files = os.listdir(WORD_DIR)
    
    tasks = []
    for _, row in df.iterrows():
        loc = str(row[col_loc]).strip()
        art = str(row[col_art]).strip()
        alb = str(row[col_alb]).strip()

        if not alb or alb == "nan":
            continue

        # 이미 워드 문서가 존재하는지 확인 (위치 번호 기준)
        already_exists = any(f"위치_{loc}_" in f or f"_{loc}_" in f or f.startswith(f"{loc}_") for f in existing_files)
        if not already_exists:
            tasks.append((loc, art, alb))

    print(f"📌 총 {len(df)}개 목록 중 새로 작성할 음반: {len(tasks)}개\n")

    for i, (loc, art, alb) in enumerate(tasks, 1):
        print(f"[{i}/{len(tasks)}] AI 검색 및 작성 중 -> [{loc}] {art} - {alb}")
        ai_result = generate_lp_content(loc, art, alb)
        if ai_result:
            create_word_doc(loc, art, alb, ai_result)
        time.sleep(1.5)  # API 안정성을 위한 텀

    # 갱신된 내용을 음악비서 DB에 반영하도록 캐시 삭제
    db_cache = r"D:\LP_Agent\lp_database.json"
    if os.path.exists(db_cache):
        os.remove(db_cache)

    print("\n🎉 모든 신규 음반의 워드 문서 작성이 완료되었습니다!")
    input("\n엔터(Enter) 키를 누르면 종료됩니다...")

if __name__ == "__main__":
    main()