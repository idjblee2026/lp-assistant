import os
import re
from docx import Document
from docx2pdf import convert
from pypdf import PdfWriter

source_dir = r"D:\LP_Agent\word_reports"
temp_pdf_dir = r"D:\LP_Agent\pdf_temp_at"

os.makedirs(temp_pdf_dir, exist_ok=True)

files = sorted([f for f in os.listdir(source_dir) if f.endswith(".docx") and not f.startswith("~$")])

print(f"총 {len(files)}개 Word 문서에 '@' 태그(@0001 / @701)를 적용합니다...")

for filename in files:
    doc_path = os.path.join(source_dir, filename)
    doc = Document(doc_path)
    
    # 1. 번호 : 0001 -> 번호 : @0001
    def replace_num(match):
        num = int(match.group(1))
        return f"번호 : @{num:04d}"

    # 2. 위치 : 701 -> 위치 : @701
    def replace_loc(match):
        num = int(match.group(1))
        return f"위치 : @{num}"

    for p in doc.paragraphs:
        if "번호" in p.text:
            p.text = re.sub(r"번호\s*:\s*@?(\d+)", replace_num, p.text)
        if "위치" in p.text:
            p.text = re.sub(r"위치\s*:\s*@?(\d+)", replace_loc, p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "번호" in p.text:
                        p.text = re.sub(r"번호\s*:\s*@?(\d+)", replace_num, p.text)
                    if "위치" in p.text:
                        p.text = re.sub(r"위치\s*:\s*@?(\d+)", replace_loc, p.text)

    doc.save(doc_path)

print("[1단계 완료] Word 문서 내 '@' 태그 적용 완료")

# 3. PDF 변환 및 병합
print("[2단계 진행] PDF 일괄 생성 및 통합 검색본 병합 중...")
convert(source_dir, temp_pdf_dir)

merger = PdfWriter()
pdf_files = sorted([os.path.join(temp_pdf_dir, f) for f in os.listdir(temp_pdf_dir) if f.endswith(".pdf")])

for pdf in pdf_files:
    merger.append(pdf)

final_pdf_path = r"D:\LP_Agent\LP_전체음반_통합검색본.pdf"
with open(final_pdf_path, "wb") as f_out:
    merger.write(f_out)
merger.close()

# 임시 PDF 파일 삭제
for f in pdf_files:
    try:
        os.remove(f)
    except:
        pass
try:
    os.rmdir(temp_pdf_dir)
except:
    pass

print("="*60)
print(f"완료되었습니다!")
print(f"- 검색 요령:")
print(f"  * 위치 찾기: @104 (3자리)")
print(f"  * 번호 찾기: @0104 (4자리)")
print(f"- 최종 파일: {final_pdf_path}")
print("="*60)