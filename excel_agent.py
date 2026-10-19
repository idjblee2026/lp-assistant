import os
import re
import pandas as pd
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn

# ==========================================
# 1. 경로 설정
# ==========================================
EXCEL_FILE = r"D:\LP_Agent\lp_list.xlsx"
OUTPUT_DIR = r"D:\LP_Agent\enhanced_reports"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ==========================================
# 2. 한글 폰트(맑은 고딕) 강제 적용 함수
# ==========================================
def set_korean_font(run, font_name="맑은 고딕", size_pt=10, bold=False, color=None):
    run.font.name = font_name
    run._r.get_or_add_rPr().set(qn('w:rFonts'), font_name)
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

# ==========================================
# 3. 워드 문서 작성 함수 (번호, 위치 중심)
# ==========================================
def create_lp_report(file_path, row):
    doc = docx.Document()

    # 여백 설정
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # 1) 메인 제목
    raw_num = row.get('번호', '')
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(f"LP 소장 관리 카드 - No. {raw_num}")
    set_korean_font(r_title, size_pt=18, bold=True, color=RGBColor(24, 43, 73))

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("아날로그 바이닐 컬렉션 소장 및 보관 위치 정보")
    set_korean_font(r_sub, size_pt=11, bold=False, color=RGBColor(100, 100, 100))
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2) 핵심 제원 및 위치 정보 테이블
    p_h1 = doc.add_paragraph()
    r_h1 = p_h1.add_run("■ 소장 관리 및 수납 위치")
    set_korean_font(r_h1, size_pt=13, bold=True, color=RGBColor(31, 78, 121))

    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'

    # 엑셀의 모든 컬럼 값을 깔끔하게 표로 정리
    items = []
    
    # 번호와 위치를 최상단에 배치
    if '번호' in row:
        items.append(("관리 번호", str(row.get('번호', '')).replace('nan', '')))
    if '위치' in row:
        items.append(("보관 위치 (수납 랙)", str(row.get('위치', '')).replace('nan', '')))
    
    # 그 외 엑셀에 존재하는 다른 컬럼들도 함께 표기
    for col in row.index:
        if col not in ['번호', '위치'] and not str(col).startswith('Unnamed'):
            val = str(row.get(col, '')).replace('nan', '')
            if val:
                items.append((str(col), val))

    for key, val in items:
        row_cells = table.add_row().cells
        # 항목명 열
        r1 = row_cells[0].paragraphs[0].add_run(key)
        set_korean_font(r1, size_pt=10, bold=True, color=RGBColor(30, 30, 30))
        # 내용 열
        r2 = row_cells[1].paragraphs[0].add_run(val if val else "-")
        set_korean_font(r2, size_pt=10, bold=False)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 3) 관리 메모 란
    p_h2 = doc.add_paragraph()
    r_h2 = p_h2.add_run("■ 청음 및 관리 메모")
    set_korean_font(r_h2, size_pt=12, bold=True, color=RGBColor(31, 78, 121))

    p_memo = doc.add_paragraph()
    p_memo.paragraph_format.line_spacing = 1.3
    r_memo = p_memo.add_run("• 음반 상태 및 클리닝 이력:\n• 청음 추천 카트리지 / 침압:\n• 특이사항:")
    set_korean_font(r_memo, size_pt=10, bold=False, color=RGBColor(80, 80, 80))

    doc.save(file_path)

# ==========================================
# 4. 일괄 실행 루프
# ==========================================
print("=" * 60)
print("엑셀 데이터 기반 LP 소장 관리 카드 일괄 생성 시작")
print("=" * 60)

try:
    df = pd.read_excel(EXCEL_FILE, sheet_name="번호인쇄")
except Exception:
    # '번호인쇄' 시트가 없을 경우 첫 번째 시트 자동 로드
    df = pd.read_excel(EXCEL_FILE, sheet_name=0)

# 컬럼명 공백 제거
df.columns = [str(c).strip() for c in df.columns]

count = 0
for idx, row in df.iterrows():
    raw_num = row.get('번호', idx + 1)
    try:
        num_str = f"{int(raw_num):03d}"
    except:
        num_str = f"{str(raw_num):>03}"

    loc = str(row.get('위치', '')).replace('nan', '').strip()
    safe_loc = re.sub(r'[\\/*?:"<>|]', '_', loc)[:15] if loc else "미지정"

    file_name = f"LP_{num_str}_위치_{safe_loc}.docx"
    file_path = os.path.join(OUTPUT_DIR, file_name)

    create_lp_report(file_path, row)
    count += 1
    print(f"[{count}/{len(df)}] No.{num_str} (위치: {loc}) 관리 카드 저장 완료")

print("=" * 60)
print(f"총 {count}개의 워드 관리 카드가 'enhanced_reports' 폴더에 정상 생성되었습니다.")
print("=" * 60)